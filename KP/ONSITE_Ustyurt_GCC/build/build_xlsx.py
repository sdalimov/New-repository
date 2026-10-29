"""Заполняет RFQ Устюртского ГХК от имени ONSITE и добавляет лист 4 с расчётом количеств.
usage: python3 -I build_xlsx.py <rfq.xlsx> <out.xlsx>"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
import kp_data as D

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
ws = wb["1. Цены - Prices"]
terms = wb["2. Условия - Terms"]

# --- Шапка участника ---
ws["K6"] = D.COMPANY["name"]
ws["K7"] = D.COMPANY["address"]
ws["K8"] = D.COMPANY["contact"]
ws["K9"] = D.KP_DATE
ws["K10"] = D.VALIDITY_DAYS
ws["K11"] = D.CURRENCY
ws["K12"] = "Доставка до УГХК (местный участник) / Delivery to UGCC (local bidder)"
ws["K13"] = "Да / Yes"

# --- Лист 4: расчёт количеств ---
CS = "4. Расчёт - Calculation"
cs = wb.create_sheet(CS)
thin = Side(style="thin", color="999999")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
f = lambda **k: Font(name="Arial", size=9, **k)
yellow = PatternFill("solid", fgColor="FFF2CC")
grey = PatternFill("solid", fgColor="D9D9D9")

cs["A1"] = "РАСЧЁТ КОЛИЧЕСТВА РЕАГЕНТОВ (программа дозирования ONSITE) / QUANTITY CALCULATION"
cs["A1"].font = Font(name="Arial", size=11, bold=True)
cs["A2"] = ("Кол-во, кг = База × Дозировка (г/т или г/м³) × Множитель / 1000; предлагаемое кол-во округлено вверх до 10 кг. "
            "Жёлтые ячейки — допущения ONSITE, при изменении пересчитываются листы 4 и 1.")
cs["A2"].font = f(italic=True)
hdr = ["№", "Наименование", "База", "Ед. базы", "Дозировка, г/т (г/м³)", "Множитель",
       "Пояснение к множителю", "Расчётное кол-во, кг", "Предлагаемое кол-во, кг", "Допущения / комментарий"]
for i, h in enumerate(hdr, 1):
    c = cs.cell(4, i, h)
    c.font = f(bold=True); c.fill = grey; c.border = box
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
calc_row = {}
for n, c in enumerate(D.CALC, start=5):
    vals = [c["no"], c["name"], c["base"], c["base_unit"], c["dose"], c["mult"], c["mult_desc"],
            f"=C{n}*E{n}*F{n}/1000", f"=ROUNDUP(H{n}/10,0)*10", c["note"]]
    for i, v in enumerate(vals, 1):
        cell = cs.cell(n, i, v)
        cell.font = f(); cell.border = box
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i in (3, 5, 6):
            cell.fill = yellow; cell.font = f(color="0000FF")
    cs.cell(n, 3).number_format = "#,##0"
    cs.cell(n, 5).number_format = "#,##0.0##"
    cs.cell(n, 8).number_format = "#,##0.0"
    cs.cell(n, 9).number_format = "#,##0"
    cs.cell(n, 9).font = f(bold=True)
    calc_row[c["no"]] = n
for col, w in zip("ABCDEFGHIJ", [6, 40, 12, 26, 13, 10, 30, 14, 14, 60]):
    cs.column_dimensions[col].width = w
cs.freeze_panes = "C5"

# --- Лист 1: строки участника ---
link_font = Font(name="Arial", size=9, color="008000")
for r in D.items():
    row = r["row"]
    if r["qty"] == "calc":
        n = calc_row[r["no"]]
        ws[f"N{row}"] = f"='{CS}'!I{n}"
        ws[f"N{row}"].font = link_font
        label = D.CALC_BY_NO[r["no"]].get("dose_label")
        ws[f"M{row}"] = label or f"='{CS}'!E{n}&\" г/т (ppm)\""
    if r.get("trade"):
        ws[f"K{row}"] = r["trade"]
    if r.get("maker"):
        ws[f"L{row}"] = r["maker"]
    if r["price"] is not None:
        ws[f"O{row}"] = r["price"]
    if r["unit"] in ("кг", "шт."):
        ws[f"Q{row}"] = "Еврокуб 1 м³ (невозвратный)" if r["unit"] == "кг" else "Коробка, индивид. маркировка"
        ws[f"R{row}"] = D.LEAD_TIME_DAYS
        if r["unit"] == "кг":
            ws[f"S{row}"] = D.SHELF_LIFE_M
    note = D.CALC_BY_NO.get(r["no"], {}).get("note") if r["no"] == "11" else None
    if note:
        ws[f"T{row}"] = note

# --- Лист 2: подтверждение требований ---
for row in range(7, 43):
    if isinstance(terms[f"A{row}"].value, int):
        terms[f"E{row}"] = "Да / Yes"

wb.active = 0
wb.save(out)
print("saved", out)
