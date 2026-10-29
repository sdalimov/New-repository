"""КП ONSITE для Устюртского ГХК в оформлении образца (КП №77 ФерганаАзот).
usage: python3 -I build_docx.py <logo.jpg> <out.docx> [--stamp stamp.jpg]"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import kp_data as D

logo, out = sys.argv[1], sys.argv[2]
stamp = sys.argv[sys.argv.index("--stamp") + 1] if "--stamp" in sys.argv else None
BLANK = "________"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.15

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin, sec.bottom_margin = Cm(1.2), Cm(1.8)
sec.header_distance, sec.footer_distance = Cm(0.6), Cm(0.6)

hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.add_run().add_picture(logo, width=Cm(4.3))

for i, line in enumerate((D.COMPANY["footer1"], D.COMPANY["footer2"])):
    fp = sec.footer.paragraphs[0] if i == 0 else sec.footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(line); r.font.name = "Calibri"; r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0xA6, 0xA6, 0xA6)


def para(text="", bold=False, align=None, indent=None, after=0, before=0, blank=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.left_indent = indent
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if text:
        r = p.add_run(text); r.bold = bold
    if blank:
        r = p.add_run(BLANK); r.bold = bold; r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def runs(p, parts):
    for text, bold in parts:
        p.add_run(text).bold = bold


para(f"Дата: {D.KP_DATE}", indent=Cm(1.2))
para("№" + (str(D.KP_NO) if D.KP_NO else ""), indent=Cm(1.2), blank=not D.KP_NO, after=12)
for line in ("Руководителю", "СП ООО «Uz-Kor Gas Chemical»", "(Устюртский ГХК)"):
    para(line, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
para("Коммерческое предложение № " + (str(D.KP_NO) if D.KP_NO else ""), bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=8, blank=not D.KP_NO)
para("Уважаемые господа!", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

body = [
    "Компания ONSITE, основанная в 2019 году, занимается производством химических продуктов для разных "
    "сфер промышленности. Деятельность нашей компании фокусирована на локализации импортных товаров в "
    "Узбекистане. На данный момент мы имеем большое количество клиентов в Узбекистане и за его пределами. "
    "Среди них АО НГМК, СП ООО «СамАвто», Павлодарский Алюминиевый Завод (Казахстан), Sayhun Bottlers, "
    "Akfa Building Materials, Knauf и другие.",
    "В ответ на Ваш запрос ценового предложения направляем Вам коммерческое предложение на поставку "
    "химических реагентов и сопутствующих услуг для установки очистки сырой воды, системы оборотной воды "
    "(градирня), установки подготовки котловой питательной воды и установки очистки сточных вод для "
    "повторного использования на 2027 год, в том числе антискаланта для мембран обратного осмоса "
    "собственного производства – CLEAN RO-1.",
]
for t in body:
    p = para(t, after=4); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)

# --- Таблица ---
widths = [Cm(1.1), Cm(5.4), Cm(2.2), Cm(1.4), Cm(1.9), Cm(2.6), Cm(3.1)]
hdr = ["#", "Наименование", "Упак.", "Ед.", "Кол-во", "Цена, сум", "ИТОГО, сум"]
tbl = doc.add_table(rows=0, cols=7)
tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color)
    tcPr.append(sh)


def fill(cells, values, bold=False, size=10, right=(), hl=()):
    for i, (c, v) in enumerate(zip(cells, values)):
        c.width = widths[i] if len(cells) == 7 else c.width
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = c.paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        if i in right:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(v); r.bold = bold; r.font.size = Pt(size)
        if i in hl:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def money(x):
    return f"{x:,.0f}"


fill(tbl.add_row().cells, hdr, bold=True, size=10.5)
total, missing, prefixed = 0, False, False
for title, rows in D.SECTIONS:
    row = tbl.add_row(); m = row.cells[0].merge(row.cells[6])
    fill([m], [title], bold=True, size=10); shade(m, "F2F2F2")
    for r in rows:
        if "group" in r:
            row = tbl.add_row(); m = row.cells[1].merge(row.cells[6])
            fill([row.cells[0], m], ["", r["group"]], bold=True, size=10)
            continue
        q = D.qty_of(r)
        if r["price"] is None:
            missing = True; price, amount = BLANK, BLANK
        else:
            pre = r.get("prefix", ""); prefixed = prefixed or bool(pre)
            total += q * r["price"]; price, amount = pre + money(r["price"]), pre + money(q * r["price"])
        hl = (5, 6) if r["price"] is None else ()
        fill(tbl.add_row().cells, [r["no"] + ".", r["name"], r["pack"], r["unit"], money(q), price, amount],
             right=(4, 5, 6), hl=hl)
fill(tbl.add_row().cells, ["", "ИТОГО:", "", "", "", "", BLANK if missing else ("от " if prefixed else "") + money(total)],
     bold=True, right=(6,), hl=(6,) if missing else ())
for row in tbl.rows:
    for i, c in enumerate(row.cells):
        c.width = widths[i]
trPr = tbl.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
for row in tbl.rows:
    cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); row._tr.get_or_add_trPr().insert(0, cs)
for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
    gc.set(qn("w:w"), str(int(w.twips)))
tblPr = tbl._tbl.tblPr
lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
tw = tblPr.find(qn("w:tblW"))
if tw is None:
    tw = OxmlElement("w:tblW"); tblPr.append(tw)
tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(sum(w.twips for w in widths))))

para("", after=6)
conds = [
    ("Цены: ", f"в сумах, {D.VAT_TEXT}, включают упаковку (еврокубы 1 м³, невозвратные), маркировку и доставку"),
    ("Условия оплаты: ", "100% по факту поставки товара"),
    ("Условия поставки: ", "доставка до УГХК, Кырккыз, Кунградский р-н, Республика Каракалпакстан"),
    ("Срок поставки: ", f"до {D.LEAD_TIME_DAYS} дней после заключения контракта, партиями по согласованному графику"),
    ("Срок годности: ", f"{D.SHELF_LIFE_M} месяца с даты производства; остаточный срок при поставке — не менее 75%"),
    ("Количество: ", "рассчитано по программе дозирования ONSITE на основании исходных данных ТЗ "
                     "(расчёт — лист 4 приложенной формы запроса)"),
    ("Апробация: ", "проводится на площадке Заказчика с участием 1 представителя ONSITE в течение до 3 месяцев"),
    ("Срок действия предложения: ", f"{D.VALIDITY_DAYS} дней"),
]
for k, v in conds:
    p = para(); runs(p, [(k, False), (v, False)])
para("", after=8)

sig = ["С Уважением,", "Директор ООО “ONSITE”", D.COMPANY["director"], "Тел.: " + D.COMPANY["phone"]]
if stamp:
    t = doc.add_table(rows=1, cols=2); t.autofit = False
    left, right = t.rows[0].cells
    left.width, right.width = Cm(6.5), Cm(6)
    lp = left.paragraphs[0]
    for i, s in enumerate(sig):
        p = lp if i == 0 else left.add_paragraph()
        p.add_run(s).bold = True
    em = left.add_paragraph().add_run(D.COMPANY["email"]); em.bold = True; em.underline = True
    em.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    right.paragraphs[0].add_run().add_picture(stamp, width=Cm(4.5))
else:
    for s in sig:
        para(s, bold=True)
    em = para().add_run(D.COMPANY["email"]); em.bold = True; em.underline = True
    em.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)

doc.save(out)
print("saved", out, "| prices missing:", missing)
